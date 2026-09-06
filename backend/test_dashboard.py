from __future__ import annotations

import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook

from . import app_trace
from .dashboard_metrics import canonical_route_key, parse_measurement_workbook
from .dashboard_storage import (
    DashboardRouteNotFoundError,
    DashboardStorageConfigurationError,
    DashboardStorageOperationError,
    RouteDashboardStorage,
    aggregate_route_series,
)


def snapshot(
    month: str,
    *,
    loss: float | None = None,
    position: float | None = None,
    dkd: float | None = None,
    required: float | None = None,
    assessment: str | None = None,
    sheet: str = "TYN - CPA",
) -> dict:
    return {
        "region": "DBB",
        "route_key": "cpa--tyn",
        "route_name": sheet,
        "measurement_month": f"{month}-01",
        "worst_event_loss_db": loss,
        "worst_event_position_km": position,
        "utilization_percent": 25.0 if dkd is not None else None,
        "dkd_core_percent": dkd,
        "dkd_required_percent": required,
        "assessment": assessment,
        "source_file": f"DBB {month}.xlsx",
        "source_sheet": sheet,
        "source_format": "stv",
        "source_sha256": month,
        "warnings": [],
    }


class DashboardMetricTests(unittest.TestCase):
    def test_route_key_merges_direction_but_preserves_suffix(self) -> None:
        self.assertEqual(canonical_route_key("TYN - CPA"), "cpa--tyn")
        self.assertEqual(canonical_route_key("CPA–TYN"), "cpa--tyn")
        self.assertNotEqual(
            canonical_route_key("TYN - CPA B"),
            canonical_route_key("TYN - CPA"),
        )

    def test_parser_reads_stv_event_and_qd48_values_without_recalculation(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "DBB T7.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "CPA - TYN"
            headers = [
                "Trạng thái đồ thị", "Tệp", "Định dạng", "Sợi", "Bước sóng",
                "Suy hao tổng, dB", "Chiều dài, km", "Suy hao TB, dB/km", "Đánh giá",
            ]
            for column, label in enumerate(headers, 1):
                sheet.cell(5, column, label)
            sheet.cell(5, 10, 2.5)
            sheet.cell(5, 11, 40.09)
            sheet.cell(6, 2, "1.sor")
            sheet.cell(6, 9, "Không đạt")
            sheet.cell(6, 10, 1.2)
            sheet.cell(6, 11, 4.57)
            for column, label in enumerate(
                ["Tỉ lệ khai thác", "DKD Core", "DKD yêu cầu", "Đánh giá"], 1
            ):
                sheet.cell(10, column, label)
            sheet.cell(11, 1, 0.25)
            sheet.cell(11, 2, 0.29166667)
            sheet.cell(11, 3, 70)
            sheet.cell(11, 4, "Không đạt")
            workbook.save(path)

            rows = parse_measurement_workbook(
                path, region="DBB", measurement_month="2026-07"
            )

        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row.route_key, "cpa--tyn")
        self.assertEqual(row.worst_event_loss_db, 4.57)
        self.assertEqual(row.worst_event_position_km, 40.09)
        self.assertEqual(row.utilization_percent, 25.0)
        self.assertAlmostEqual(row.dkd_core_percent or 0, 29.166667)
        self.assertEqual(row.dkd_required_percent, 70.0)
        self.assertEqual(row.assessment, "Không đạt")

    def test_six_month_series_keeps_gaps_and_uses_monthly_maximum(self) -> None:
        rows = [
            snapshot("2026-03", loss=18.779, position=2.54),
            snapshot("2026-04", loss=20.0, position=4.0, dkd=50, required=80, assessment="Không đạt"),
            snapshot("2026-04", loss=20.708, position=2.456),
            snapshot("2026-07", loss=4.57, position=40.09, dkd=72, required=70, assessment="Đạt"),
        ]

        result = aggregate_route_series(
            rows, region="DBB", route_key="cpa--tyn", months=6
        )

        self.assertEqual(result["period_start"], "2026-02-01")
        self.assertEqual(result["period_end"], "2026-07-01")
        self.assertIsNone(result["months"][0]["worst_event_loss_db"])
        self.assertEqual(result["months"][2]["worst_event_loss_db"], 20.708)
        self.assertEqual(result["months"][2]["dkd_core_percent"], 50.0)
        self.assertEqual(result["months"][5]["assessment"], "Đạt")

    def test_local_storage_never_mixes_regions(self) -> None:
        rows = [snapshot("2026-07", loss=4.57)]
        rows.append({**snapshot("2026-07", loss=99), "region": "TBB"})
        with TemporaryDirectory() as directory:
            path = Path(directory) / "dashboard.json"
            path.write_text(json.dumps({"snapshots": rows}), encoding="utf-8")
            storage = RouteDashboardStorage(local_data_file=path)
            routes = storage.list_routes("DBB")
            result = storage.route_series("DBB", routes[0]["route_key"])
        self.assertEqual(len(routes), 1)
        self.assertEqual(result["months"][-1]["worst_event_loss_db"], 4.57)


class DashboardEndpointTests(unittest.TestCase):
    @staticmethod
    def payload(response: object) -> dict:
        return json.loads(response.body.decode("utf-8"))

    def test_routes_endpoint_returns_storage_data(self) -> None:
        storage = Mock()
        storage.list_routes.return_value = [{"route_key": "cpa--tyn", "route_name": "TYN - CPA"}]
        with patch.object(app_trace, "get_dashboard_storage", return_value=storage):
            response = app_trace.get_dashboard_routes("DBB")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.payload(response)["data"][0]["route_key"], "cpa--tyn")

    def test_route_endpoint_maps_not_found_and_configuration_errors(self) -> None:
        missing_route = Mock()
        missing_route.route_series.side_effect = DashboardRouteNotFoundError("missing")
        with patch.object(app_trace, "get_dashboard_storage", return_value=missing_route):
            response = app_trace.get_dashboard_route("DBB", "cpa--tyn")
        self.assertEqual(response.status_code, 404)

        with patch.object(
            app_trace,
            "get_dashboard_storage",
            side_effect=DashboardStorageConfigurationError("secret"),
        ):
            response = app_trace.get_dashboard_routes("DBB")
        self.assertEqual(response.status_code, 503)
        self.assertNotIn("secret", self.payload(response)["detail"])

    def test_route_endpoint_maps_operation_error_without_leaking_details(self) -> None:
        storage = Mock()
        storage.route_series.side_effect = DashboardStorageOperationError("network detail")
        with patch.object(app_trace, "get_dashboard_storage", return_value=storage):
            response = app_trace.get_dashboard_route("DBB", "cpa--tyn")
        self.assertEqual(response.status_code, 502)
        self.assertNotIn("network detail", self.payload(response)["detail"])


if __name__ == "__main__":
    unittest.main()
