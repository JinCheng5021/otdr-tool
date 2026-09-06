from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
import hashlib
import math
from pathlib import Path
import re
import unicodedata
from typing import Any, Iterable

from openpyxl import load_workbook


REGIONS = ("B-N", "TBB", "DBB")
ASSESSMENTS = {"dat": "Đạt", "khong dat": "Không đạt"}


def fold_text(value: object) -> str:
    text = str(value or "").strip().lower().replace("đ", "d")
    normalized = "".join(
        char
        for char in unicodedata.normalize("NFD", text)
        if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"\s+", " ", normalized).strip()


def normalize_region(value: object) -> str:
    folded = fold_text(value).replace("_", "-")
    compact = re.sub(r"[^a-z0-9]", "", folded)
    if compact == "dbb":
        return "DBB"
    if compact == "tbb":
        return "TBB"
    if compact == "bn":
        return "B-N"
    raise ValueError("region must be one of B-N, TBB or DBB")


def display_route_name(value: object) -> str:
    text = unicodedata.normalize("NFC", str(value or "")).strip()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s*[–—]\s*", " - ", text)
    text = re.sub(r"\s*-\s*", " - ", text)
    return text.strip(" -")


def canonical_route_key(value: object) -> str:
    """Build a direction-independent key without removing physical suffixes."""
    display = display_route_name(value)
    parts = [part.strip() for part in re.split(r"\s+-\s+", display) if part.strip()]
    folded_parts = [re.sub(r"[^a-z0-9]+", "-", fold_text(part)).strip("-") for part in parts]
    if len(folded_parts) == 2 and all(folded_parts):
        return "--".join(sorted(folded_parts))
    return re.sub(r"[^a-z0-9]+", "-", fold_text(display)).strip("-")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _finite_number(value: object) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    number = float(value)
    return number if math.isfinite(number) else None


def _percent(value: object) -> float | None:
    number = _finite_number(value)
    if number is None:
        return None
    if abs(number) <= 1.0000001:
        number *= 100.0
    return round(number, 6)


def _assessment(value: object) -> str | None:
    return ASSESSMENTS.get(fold_text(value))


def _month_start(value: str | date | datetime) -> str:
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        text = str(value).strip()
        match = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-(\d{1,2}))?", text)
        if not match:
            raise ValueError("measurement month must use YYYY-MM or YYYY-MM-DD")
        parsed = date(int(match.group(1)), int(match.group(2)), int(match.group(3) or 1))
    return parsed.replace(day=1).isoformat()


@dataclass
class DashboardSnapshot:
    region: str
    route_key: str
    route_name: str
    measurement_month: str
    worst_event_loss_db: float | None = None
    worst_event_position_km: float | None = None
    utilization_percent: float | None = None
    dkd_core_percent: float | None = None
    dkd_required_percent: float | None = None
    assessment: str | None = None
    source_file: str = ""
    source_sheet: str = ""
    source_format: str = "unknown"
    source_sha256: str = ""
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _find_main_header(ws: Any) -> tuple[int | None, dict[str, int]]:
    aliases = {
        "file": {"file", "tep"},
        "fiber": {"fiber", "soi"},
        "wavelength": {"wavelength", "buoc song"},
        "loss": {"loss, db", "suy hao tong, db"},
        "length": {"length, km", "chieu dai, km"},
        "assessment": {"danh gia"},
    }
    for row in range(1, min(ws.max_row, 25) + 1):
        columns: dict[str, int] = {}
        for column in range(1, min(ws.max_column, 40) + 1):
            label = fold_text(ws.cell(row, column).value)
            for canonical, choices in aliases.items():
                if label in choices:
                    columns[canonical] = column
        if set(aliases).issubset(columns):
            return row, columns
    return None, {}


def _find_summary_table(ws: Any) -> tuple[int | None, dict[str, int]]:
    required = {"ti le khai thac", "dkd core", "dkd yeu cau", "danh gia"}
    for row in range(1, ws.max_row + 1):
        labels = {
            fold_text(ws.cell(row, column).value): column
            for column in range(1, min(ws.max_column, 16) + 1)
            if ws.cell(row, column).value is not None
        }
        if required.issubset(labels):
            return row, labels
    return None, {}


def _find_alternative_assessment(formula_ws: Any, value_ws: Any) -> tuple[Any, Any]:
    for row in range(1, min(formula_ws.max_row, 16) + 1):
        for column in range(1, min(formula_ws.max_column, 30) + 1):
            if fold_text(formula_ws.cell(row, column).value) != "do kha dung core":
                continue
            score = value_ws.cell(row, column + 1).value
            assessment = value_ws.cell(row, column + 2).value
            if _assessment(assessment) is not None:
                return score, assessment
    return None, None


def _find_raw_matrix(ws: Any) -> tuple[int | None, list[tuple[int, float]]]:
    candidates: list[tuple[tuple[int, int, int], int, list[tuple[int, float]]]] = []
    for row in range(1, min(ws.max_row, 14) + 1):
        first_label = fold_text(ws.cell(row, 1).value)
        if first_label not in {"soi", "km/soi", ""}:
            continue
        positions = [
            (column, number)
            for column in range(2, ws.max_column + 1)
            if (number := _finite_number(ws.cell(row, column).value)) is not None
            and number >= 0
        ]
        fibers = sum(
            _finite_number(ws.cell(data_row, 1).value) is not None
            for data_row in range(row + 1, min(ws.max_row, row + 10) + 1)
        )
        if len(positions) >= 2 and fibers >= 3:
            candidates.append(((int(first_label in {"soi", "km/soi"}), fibers, len(positions)), row, positions))
    if not candidates:
        return None, []
    _, row, positions = max(candidates, key=lambda item: item[0])
    return row, positions


def _best_event_from_table(
    formula_ws: Any,
    value_ws: Any,
    header_row: int,
    assessment_column: int,
    data_end: int,
) -> tuple[float | None, float | None, Counter[str]]:
    statuses: Counter[str] = Counter()
    for row in range(header_row + 1, data_end + 1):
        status = _assessment(value_ws.cell(row, assessment_column).value)
        if status:
            statuses[status] += 1

    best_loss: float | None = None
    best_position: float | None = None
    for column in range(assessment_column + 1, formula_ws.max_column + 1):
        position = _finite_number(value_ws.cell(header_row, column).value)
        if position is None:
            position = _finite_number(formula_ws.cell(header_row, column).value)
        if position is None:
            continue
        for row in range(header_row + 1, data_end + 1):
            loss = _finite_number(value_ws.cell(row, column).value)
            if loss is None:
                loss = _finite_number(formula_ws.cell(row, column).value)
            if loss is not None and loss > 0 and (best_loss is None or loss > best_loss):
                best_loss, best_position = loss, position
    return best_loss, best_position, statuses


def _best_event_from_raw(
    formula_ws: Any,
    value_ws: Any,
    header_row: int,
    positions: Iterable[tuple[int, float]],
) -> tuple[float | None, float | None]:
    best_loss: float | None = None
    best_position: float | None = None
    for row in range(header_row + 1, formula_ws.max_row + 1):
        fiber = _finite_number(value_ws.cell(row, 1).value)
        if fiber is None:
            fiber = _finite_number(formula_ws.cell(row, 1).value)
        if fiber is None:
            continue
        for column, position in positions:
            loss = _finite_number(value_ws.cell(row, column).value)
            if loss is None:
                loss = _finite_number(formula_ws.cell(row, column).value)
            if loss is not None and loss > 0 and (best_loss is None or loss > best_loss):
                best_loss, best_position = loss, position
    if best_position is not None and best_position > 500:
        best_position /= 1000.0
    return best_loss, best_position


def parse_measurement_workbook(
    path: str | Path,
    *,
    region: str,
    measurement_month: str | date | datetime,
) -> list[DashboardSnapshot]:
    source = Path(path)
    normalized_region = normalize_region(region)
    normalized_month = _month_start(measurement_month)
    digest = sha256_file(source)
    formula_book = load_workbook(source, data_only=False, read_only=False)
    value_book = load_workbook(source, data_only=True, read_only=False)
    snapshots: list[DashboardSnapshot] = []
    try:
        for formula_ws in formula_book.worksheets:
            value_ws = value_book[formula_ws.title]
            route_name = display_route_name(formula_ws.title)
            route_key = canonical_route_key(route_name)
            if not route_key:
                continue

            header_row, columns = _find_main_header(formula_ws)
            summary_row, summary_columns = _find_summary_table(formula_ws)
            utilization = dkd_core = dkd_required = None
            assessment = None
            warnings: list[str] = []
            if summary_row is not None:
                result_row = summary_row + 1
                utilization = _percent(value_ws.cell(result_row, summary_columns["ti le khai thac"]).value)
                dkd_core = _percent(value_ws.cell(result_row, summary_columns["dkd core"]).value)
                dkd_required = _percent(value_ws.cell(result_row, summary_columns["dkd yeu cau"]).value)
                assessment = _assessment(value_ws.cell(result_row, summary_columns["danh gia"]).value)
                if assessment is None:
                    assessment = _assessment(formula_ws.cell(result_row, summary_columns["danh gia"]).value)
            else:
                score, legacy_assessment = _find_alternative_assessment(formula_ws, value_ws)
                dkd_core = _percent(score)
                assessment = _assessment(legacy_assessment)
                if assessment is not None:
                    warnings.append("legacy_qd48_summary")
                else:
                    warnings.append("missing_qd48_summary")

            worst_loss = worst_position = None
            source_format = "unknown"
            if header_row is not None:
                data_end = summary_row - 2 if summary_row else formula_ws.max_row
                worst_loss, worst_position, _ = _best_event_from_table(
                    formula_ws,
                    value_ws,
                    header_row,
                    columns["assessment"],
                    data_end,
                )
                source_format = (
                    "stv"
                    if fold_text(formula_ws.cell(header_row, columns["file"]).value) == "tep"
                    else "fast"
                )
            else:
                raw_row, positions = _find_raw_matrix(value_ws)
                if raw_row is None:
                    raw_row, positions = _find_raw_matrix(formula_ws)
                if raw_row is not None:
                    worst_loss, worst_position = _best_event_from_raw(
                        formula_ws, value_ws, raw_row, positions
                    )
                    source_format = "raw_matrix"
                else:
                    warnings.append("missing_measurement_matrix")

            if worst_loss is None and assessment is None and dkd_core is None:
                continue
            snapshots.append(
                DashboardSnapshot(
                    region=normalized_region,
                    route_key=route_key,
                    route_name=route_name,
                    measurement_month=normalized_month,
                    worst_event_loss_db=round(worst_loss, 6) if worst_loss is not None else None,
                    worst_event_position_km=round(worst_position, 6) if worst_position is not None else None,
                    utilization_percent=utilization,
                    dkd_core_percent=dkd_core,
                    dkd_required_percent=dkd_required,
                    assessment=assessment,
                    source_file=source.name,
                    source_sheet=formula_ws.title,
                    source_format=source_format,
                    source_sha256=digest,
                    warnings=warnings,
                )
            )
    finally:
        formula_book.close()
        value_book.close()
    return snapshots
